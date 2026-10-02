"""A QC batch on disk: QC's review sheet and the transcripts it labels.

    data/qc/<batch>/
      *.xlsx                   QC's review sheet (exactly one)
      *.zip | transcripts/     the transcripts, as QC sent them
      inputs/                  written here: one JSON per call, metadata folded in
      runs/<run id>/           written by `qc_eval_report`

The sheet's layout changes between batches, so columns are found by name:
the header row is the first row (of any sheet, within its first five rows)
that has `Call ID` *and* a label column. Rows join transcripts on **Call ID**
— the trailing integer of the file name: the sheet masks the phone number,
and its date column was read dd/mm as mm/dd.
"""
from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path
from typing import Dict, Iterator, Optional, Tuple

CRITERIA = "thái độ đtv"
#: QC's verdict, by the names batches have used for it — first found wins.
LABEL = ("kq qc hiện tại", "kq thủ công", "qc result")
COMMENT = ("nhận xét của qc cho ai", "nhận xét của qc", "qc comment", "dẫn chứng thủ công")
CALL_ID, CALL_CODE, CLOSED_BY, CRITERION = "call id", "call code", "bên ngắt máy", "tiêu chí"
CLEAN = "tích cực"

_TAIL_ID = re.compile(r"(\d+)(?:\.wav|\.json)?$")


def call_id_of(name: str) -> Optional[str]:
    """The trailing integer of a transcript's file name — the join key."""
    m = _TAIL_ID.search(Path(name).name)
    return m.group(1) if m else None


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).lstrip("﻿").strip()


def read_review(xlsx: Path) -> Dict[str, dict]:
    """`{call_id: {call_code, closed_by, label, violation, comment}}` for the
    `Thái độ ĐTV` rows. `label` is QC's verdict as written, `violation` is
    it read (anything but `Tích cực`), None where QC left it blank."""
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        for _ in range(5):
            header = next(rows, None)
            if header is None:
                break
            names = [_cell(c).casefold() for c in header]
            label = next((n for n in LABEL if n in names), None)
            if CALL_ID not in names or label is None:
                continue
            col = {n: names.index(n) for n in names if n}
            comment = next((n for n in COMMENT if n in names), None)
            out: Dict[str, dict] = {}
            for r in rows:
                get = lambda n: _cell(r[col[n]]) if n in col and col[n] < len(r) else ""  # noqa: E731
                cid = get(CALL_ID)
                if not cid or (CRITERION in col and get(CRITERION).casefold() != CRITERIA):
                    continue
                verdict = get(label)
                out[cid] = {
                    "call_code": get(CALL_CODE),
                    "closed_by": get(CLOSED_BY),
                    "label": verdict or None,
                    "violation": (verdict.casefold() != CLEAN) if verdict else None,
                    "comment": get(comment) if comment else "",
                }
            return out
    raise ValueError(f"{xlsx.name}: no sheet has a header row with 'Call ID' and one of {list(LABEL)}")


def _transcripts(batch: Path) -> Iterator[Tuple[str, bytes]]:
    """(file name, content) of every transcript: a zip's entries, or a
    `transcripts/` folder's files — `.wav` names holding JSON, as QC sends them."""
    folder = batch / "transcripts"
    if folder.is_dir():
        for p in sorted(folder.iterdir()):
            if p.suffix in (".wav", ".json"):
                yield p.name, p.read_bytes()
        return
    zips = sorted(batch.glob("*.zip"))
    if len(zips) != 1:
        raise FileNotFoundError(f"{batch}: put the transcripts in transcripts/ or as one .zip (found {len(zips)})")
    with zipfile.ZipFile(zips[0]) as zf:
        for info in zf.infolist():
            if not info.is_dir():
                yield Path(info.filename).name, zf.read(info)


def review_of(batch: Path) -> Path:
    sheets = sorted(batch.glob("*.xlsx"))
    if len(sheets) != 1:
        raise FileNotFoundError(f"{batch}: expected one QC review .xlsx, found {len(sheets)}")
    return sheets[0]


def prepare(batch: Path) -> dict:
    """Write `inputs/<stem>.json` for every transcript QC labelled, with
    `call_code` and `closed_by` from the sheet (the transcripts carry
    neither, and a file without `call_code` would be skipped). Returns the
    reviewed rows and what did not join, for the report to print."""
    review = read_review(review_of(batch))
    inputs = batch / "inputs"
    inputs.mkdir(exist_ok=True)
    seen, no_row, unlabelled = set(), [], []
    for name, raw in _transcripts(batch):
        cid = call_id_of(name)
        row = review.get(cid or "")
        if row is None:
            no_row.append(name)
            continue
        seen.add(cid)
        if row["violation"] is None:
            unlabelled.append(name)
            continue
        doc = json.loads(raw.decode("utf-8"))
        meta = doc.setdefault("metadata", {})
        if row["call_code"]:
            meta["call_code"] = row["call_code"]
        # `closed_by` is read as `or "YES"`: set it only when the sheet has one.
        if row["closed_by"]:
            meta["closed_by"] = row["closed_by"]
        stem = Path(name).stem
        (inputs / f"{stem}.json").write_text(json.dumps(doc, ensure_ascii=False, indent=2), encoding="utf-8")
        row["file"] = stem
    return {"review": review, "no_row": no_row, "unlabelled": unlabelled,
            "no_transcript": sorted(set(review) - seen)}
