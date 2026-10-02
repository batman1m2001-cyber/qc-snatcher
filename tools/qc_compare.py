"""Our sentiment verdicts against QC's, as a cross-tab.

    uv run python -m tools.qc_compare --scored outputs/<batch> --xlsx <review>.xlsx \
        --out outputs/<batch>_vs_qc.xlsx

Joins each scored call to its QC row on **Call ID** — the trailing integer
of the file name; the xlsx masks the phone and its date column is not
reliable. Cross-tabs QC's manual verdict (`KQ thủ công`) and QC's earlier
automatic one (`KQ Tự động`) against our `Result`, and writes one row per
call plus the tables.

A review sheet that was never filled in looks like a returned one: if the
manual column is empty everywhere, this says so — there is no ground truth
to compare against, and a table would only look authoritative.

JSON + xlsx in, xlsx out. No LLM, no network.
"""
from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any, Dict, List, Optional

CRITERIA = "Thái độ ĐTV"
CALL_ID = "call id"
MANUAL = "kq thủ công"
AUTO = "kq tự động"


def call_id_of(stem: str) -> Optional[str]:
    """The trailing integer of a file stem — the join key with QC's sheet."""
    m = re.search(r"(\d+)$", stem)
    return m.group(1) if m else None


def our_verdicts(scored: Path) -> Dict[str, Dict[str, str]]:
    """Call ID → {file, result} from a folder of scored outputs."""
    out = {}
    for p in sorted(scored.glob("*.json")):
        cid = call_id_of(p.stem)
        if cid is None:
            continue
        data = json.loads(p.read_text(encoding="utf-8"))
        row = next((r for r in data.get("Sentiment") or [] if r.get("Criteria") == CRITERIA), None)
        out[cid] = {"file": p.name, "result": (row or {}).get("Result") or ("error" if "error" in data else "")}
    return out


def qc_rows(xlsx: Path) -> Dict[str, Dict[str, str]]:
    """Call ID → {manual, auto} from the first sheet whose header row names a Call ID."""
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        for header in rows:
            names = [str(c or "").strip().lower() for c in header]
            if CALL_ID not in names:
                continue
            col = {n: names.index(n) for n in (CALL_ID, MANUAL, AUTO) if n in names}
            out = {}
            for r in rows:
                raw = r[col[CALL_ID]] if col[CALL_ID] < len(r) else None
                if raw in (None, ""):
                    continue
                cid = str(int(raw)) if isinstance(raw, (int, float)) else str(raw).strip()
                out[cid] = {k: str(r[col[c]] or "").strip() if c in col and col[c] < len(r) else ""
                            for k, c in (("manual", MANUAL), ("auto", AUTO))}
            return out
    raise ValueError(f"{xlsx.name}: no sheet has a 'Call ID' header")


def _table(pairs: List[tuple]) -> List[List[Any]]:
    counts = Counter(pairs)
    cols = sorted({b for _, b in pairs})
    return [["QC \\ ours", *cols]] + [[a, *(counts[(a, b)] for b in cols)] for a in sorted({a for a, _ in pairs})]


def compare(scored: Path, xlsx: Path, out: Path) -> Dict[str, Any]:
    ours, qc = our_verdicts(Path(scored)), qc_rows(Path(xlsx))
    joined = sorted(set(ours) & set(qc))
    rows = [[cid, ours[cid]["file"], qc[cid]["manual"], qc[cid]["auto"], ours[cid]["result"],
             (qc[cid]["manual"] == ours[cid]["result"]) if qc[cid]["manual"] else ""]
            for cid in joined]
    has_truth = any(qc[cid]["manual"] for cid in joined)

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "calls"
    ws.append(["call_id", "file", "qc_manual", "qc_auto", "ours", "agrees_with_manual"])
    for r in rows:
        ws.append(r)
    for title, idx in (("manual_vs_ours", 2), ("auto_vs_ours", 3)):
        sheet = wb.create_sheet(title)
        for line in _table([(r[idx] or "(blank)", r[4]) for r in rows]):
            sheet.append(line)
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)

    report = {"out": out, "joined": len(joined), "only_ours": len(set(ours) - set(qc)),
              "only_qc": len(set(qc) - set(ours)), "ground_truth": has_truth,
              "agree": sum(1 for r in rows if r[5] is True)}
    print(f"{out}\n  joined {report['joined']}  only ours {report['only_ours']}  only QC {report['only_qc']}")
    if not has_truth:
        print("  NO GROUND TRUTH — `KQ thủ công` is empty on every joined row. "
              "The sheet was not filled in; compare nothing against it.")
    else:
        print(f"  agrees with QC's manual verdict: {report['agree']}/{report['joined']}")
    return report


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--scored", required=True, type=Path, help="a folder of scored outputs")
    ap.add_argument("--xlsx", required=True, type=Path, help="QC's review workbook")
    ap.add_argument("--out", required=True, type=Path, help="the workbook to write")
    a = ap.parse_args(argv)
    compare(a.scored, a.xlsx, a.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
