"""Rebuild corpus.yaml from QC's review workbook.

QC reviews and re-tunes the corpus in Excel; this turns their sheet back
into the YAML the pipeline reads. Written as a script rather than done by
hand because they will send another sheet, and a hand-merge of 900 rows is
a diff nobody can check.

**The workbook is the source of truth.** Entries it does not contain are
dropped, and severities it changed are taken as given — including the ones
that flip a violation into an acceptable phrasing.

Three things it has to reconcile:

`severity` moves onto the entry
    Four tiers now — `warning`, `cao`, `nghiem_trong`, and `tich_cuc` for
    the carveout side. The old block-level `severity:` is gone, so a group
    no longer implies anything about the entries inside it.

ids
    QC's sheet lost the id on a chunk of rows. An id-less row whose content
    matches an existing entry keeps that entry's id rather than being
    treated as new, so the QC spreadsheet, the FAISS sidecar and
    `knowledge_policy` all keep pointing at the same thing.

`/` as a variant separator
    QC writes variants with `/`; the corpus uses `" | "`, because `/` is
    also an inline alternation (`anh/chị`) and splitting on it blindly
    produces fragments that match anything. See `split_slash_variants`.

    uv run python -m tools.corpus_import --dry-run         # what would change
    uv run python -m tools.corpus_import                   # write corpus.yaml (+ backup, review tsv)

Writes the corpus and a TSV of every row the slash conversion could not
resolve confidently — those are meant to be read, not trusted.
"""
from __future__ import annotations

import argparse
import collections
import re
import shutil
import sys
import unicodedata
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parents[1]
if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

XLSX = ROOT.parent / "(QC) Mô tả và Ngôn từ vi phạm_ Qc gán nhãn.xlsx"
SHEET = "tổng hợp"
CORPUS = ROOT / "knowledge/sentiment_agent/corpus.yaml"
#: The corpus as it stood before the first QC import — the reference for id
#: and group recovery, and the thing to diff against.
#:
#: NOT `corpus.yaml.bak`: that name was already taken by a pre-id-stamping
#: snapshot, it is gitignored so it never shows up in `git status`, and a
#: "skip if the backup exists" guard therefore backed up nothing at all
#: while reporting success. Using it as the reference would have been worse
#: — it has no ids, so every entry would have been re-stamped as new.
BACKUP = ROOT / "knowledge/sentiment_agent/corpus.pre-qc-import.yaml"
REVIEW = ROOT / "knowledge/sentiment_agent/qc_import_review.tsv"

#: Group for entries QC added that have no counterpart in the old corpus.
#: Dated, matching the existing `step12_20260711` convention.
NEW_GROUP = "qc_review_20260819"
NEW_GROUP_NAME = "QC review 2026-08-19 — entries added during the re-tune"

#: Excel severity text -> corpus key. `tich_cuc` marks the carveout side.
SEVERITY = {
    "warning": "warning",
    "cao": "cao",
    "nghiêm trọng": "nghiem_trong",
    "nghiem_trong": "nghiem_trong",
    "tích cực": "tich_cuc",
    "tich_cuc": "tich_cuc",
}
CARVEOUT_SEVERITY = "tich_cuc"

#: Genuine inline alternations — two ways of saying the same thing inside
#: one phrasing. Everything else touching a slash separates two phrasings.
INLINE_ALTERNATIONS = [
    r"anh\s*/\s*chị", r"chị\s*/\s*anh", r"a\s*/\s*c",
    r"em\s*/\s*tôi", r"em\s*/\s*mình", r"tôi\s*/\s*em",
    r"cố\s*tình\s*/\s*cố\s*ý", r"ông\s*/\s*bà",
]
_SENTINEL = ""

#: A converted variant shorter than this, or one that is nothing but a
#: pronoun, means the split cut through an alternation the list above does
#: not know about. Flagged rather than guessed at.
MIN_VARIANT_CHARS = 6
BARE_TOKENS = {"anh", "chị", "em", "tôi", "ac", "kh", "cb", "mình", "ngân hàng"}

#: Shortest variant distinctive enough to recover an id from. Below this a
#: phrasing is shared across entries ("ừ", "vâng") and would match the
#: wrong one.
MIN_RECOVERY_CHARS = 15


def norm(text: str) -> str:
    """Comparison key for matching QC's text against existing entries."""
    t = unicodedata.normalize("NFC", str(text or "")).lower().strip()
    return re.sub(r"\s+", " ", re.sub(r"[\"“”'`]", "", t))


def split_slash_variants(raw: str) -> tuple[list[str], list[str]]:
    """QC's `/`-separated phrasings -> variants, plus anything suspicious.

    Returns ``(variants, warnings)``. A non-empty warning list means the
    row needs a human: the split produced a fragment, which is the
    signature of having cut through an alternation.

    Slashes do three different jobs in this sheet:

    * separate whole phrasings — ``"A"/"B"`` or ``A/B`` — split
    * inline alternation — ``anh/chị`` — never split
    * shared-affix alternation — ``AC đừng la làng/ầm ĩ/quát tháo`` —
      needs the prefix redistributed, which is not something to infer, so
      it is reported instead

    The third case is why fragments are flagged rather than dropped:
    ``"ầm ĩ"`` on its own is a plausible-looking variant that matches
    almost anything, and nothing downstream would notice.
    """
    text = str(raw or "").strip()
    if not text:
        return [], []

    protected = text
    for pattern in INLINE_ALTERNATIONS:
        protected = re.sub(pattern, lambda m: m.group(0).replace("/", _SENTINEL),
                           protected, flags=re.IGNORECASE)

    # A slash between quote marks closes one phrasing and opens the next;
    # normalise it so the generic split below cannot mistake it.
    protected = re.sub(r"[\"“”]\s*/\s*[\"“”]", " | ", protected)

    parts: list[str] = []
    for chunk in re.split(r"\s*\|\s*", protected):
        parts.extend(re.split(r"\s*/\s*", chunk))

    had_slash = "/" in text
    variants, warnings = [], []
    for part in parts:
        v = part.replace(_SENTINEL, "/").strip().strip("\"“”").strip()
        if not v:
            continue
        variants.append(v)
        # Only a row that actually contained a slash can have been split by
        # one. Without this guard the check fires on genuinely short entries
        # — "xàm", "nhảm", "tổ sư" are one-word banned terms, not fragments —
        # and a review file full of false alarms is one nobody reads.
        if had_slash and (len(v) < MIN_VARIANT_CHARS or v.lower() in BARE_TOKENS):
            warnings.append(v)
    return variants, warnings


def load_existing(path: Path) -> tuple[dict, dict, dict, list]:
    """Existing corpus -> id lookup, content lookup, group names, group order."""
    raw = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    by_id, by_content, names, order = {}, {}, {}, []
    for group_id, block in raw.items():
        if not isinstance(block, dict):
            continue
        order.append(group_id)
        names[group_id] = block.get("name", "")
        for side in ("positives", "carveouts"):
            for entry in block.get(side) or []:
                if not isinstance(entry, dict):
                    continue
                content = (entry.get("content") or "").strip()
                entry_id = str(entry.get("id") or "")
                if entry_id:
                    by_id[entry_id] = {"group": group_id, "side": side,
                                       "content": content}
                if content:
                    by_content.setdefault(norm(content), entry_id)
                    # Individual variants too — QC often quotes one phrasing
                    # of a multi-variant entry. Only distinctive ones: a
                    # short variant like "ừ" appears in many entries, and
                    # matching on it hands two different QC rows the same
                    # id, which silently drops the second.
                    for v in re.split(r"\s*\|\s*", content):
                        if len(v.strip()) >= MIN_RECOVERY_CHARS:
                            by_content.setdefault(norm(v), entry_id)
    return by_id, by_content, names, order


def read_rows(xlsx: Path) -> list[dict]:
    """QC's sheet -> raw row dicts, blank rows dropped."""
    from openpyxl import load_workbook

    ws = load_workbook(xlsx, read_only=True, data_only=True)[SHEET]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not any(c is not None and str(c).strip() for c in r):
            continue
        rows.append({
            "id": str(r[0]).strip() if r[0] else "",
            "severity_raw": str(r[1]).strip() if len(r) > 1 and r[1] else "",
            "content": str(r[2]).strip() if len(r) > 2 and r[2] else "",
            "description": str(r[4]).strip() if len(r) > 4 and r[4] else "",
        })
    return rows


def build(rows: list[dict], by_id: dict, by_content: dict) -> tuple[dict, list, dict]:
    """Rows -> {group: {positives|carveouts: [entry]}}, review lines, stats."""
    next_num = {"POS": 0, "CVO": 0}
    for entry_id in by_id:
        m = re.match(r"([A-Z]+)-(\d+)$", entry_id)
        if m and m.group(1) in next_num:
            next_num[m.group(1)] = max(next_num[m.group(1)], int(m.group(2)))

    groups: dict = collections.defaultdict(lambda: {"positives": [], "carveouts": []})
    review: list[str] = []
    seen_ids: set[str] = set()
    stats = collections.Counter()

    for n, row in enumerate(rows, start=2):
        sev = SEVERITY.get(row["severity_raw"].lower())
        if sev is None:
            stats["unknown_severity"] += 1
            review.append(f"{n}\tUNKNOWN SEVERITY\t{row['severity_raw']}\t"
                          f"{row['content'][:120]}\t")
            continue

        side = "carveouts" if sev == CARVEOUT_SEVERITY else "positives"
        variants, warnings = split_slash_variants(row["content"])
        if not variants:
            stats["empty_content"] += 1
            continue
        content = " | ".join(variants)

        entry_id = row["id"]
        if entry_id and entry_id in by_id:
            stats["id_given"] += 1
        elif entry_id:
            stats["id_unknown"] += 1          # in the sheet, not in the corpus
        else:
            recovered = by_content.get(norm(row["content"])) or \
                by_content.get(norm(variants[0]))
            if recovered and recovered not in seen_ids:
                entry_id = recovered
                stats["id_recovered"] += 1
            else:
                prefix = "CVO" if side == "carveouts" else "POS"
                next_num[prefix] += 1
                entry_id = f"{prefix}-{next_num[prefix]:03d}"
                stats["id_new"] += 1

        if entry_id in seen_ids:
            stats["duplicate_id"] += 1
            review.append(f"{n}\tDUPLICATE ID\t{entry_id}\t{content[:120]}\t")
            continue
        seen_ids.add(entry_id)

        group = by_id.get(entry_id, {}).get("group") or NEW_GROUP
        entry = {"id": entry_id, "content": content, "severity": sev}
        if row["description"]:
            entry["description"] = row["description"]
        groups[group][side].append(entry)
        stats[f"side_{side}"] += 1
        stats[f"sev_{sev}"] += 1

        if warnings:
            stats["slash_review"] += 1
            review.append(f"{n}\tSLASH FRAGMENT\t{entry_id}\t"
                          f"{row['content'][:160]}\t{content[:160]}\t"
                          f"{'; '.join(warnings)}")
    return groups, review, stats


def render(groups: dict, names: dict, order: list) -> str:
    """Serialise, keeping the original group order so the diff is readable."""
    out = {}
    for group_id in order + [NEW_GROUP]:
        block = groups.get(group_id)
        if not block or not (block["positives"] or block["carveouts"]):
            continue
        rendered = {"name": names.get(group_id) or NEW_GROUP_NAME}
        for side in ("positives", "carveouts"):
            if block[side]:
                rendered[side] = block[side]
        out[group_id] = rendered
    return yaml.dump(out, allow_unicode=True, sort_keys=False,
                     default_flow_style=False, width=4096)


def main(argv=None) -> None:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", default=str(XLSX))
    ap.add_argument("--corpus", default=str(CORPUS))
    ap.add_argument("--dry-run", action="store_true",
                    help="report only; write nothing")
    a = ap.parse_args(argv)

    corpus_path = Path(a.corpus)
    # Ids and group names come from the corpus as it was BEFORE the first
    # import. Reading the live file would work only once: a second run would
    # recover ids from its own output, and any entry this import dropped
    # would then be unrecoverable.
    reference = BACKUP if BACKUP.exists() else corpus_path
    print(f"reference (ids + groups) {reference.name}")
    by_id, by_content, names, order = load_existing(reference)
    rows = read_rows(Path(a.xlsx))
    groups, review, stats = build(rows, by_id, by_content)
    text = render(groups, names, order)

    kept = stats["side_positives"] + stats["side_carveouts"]
    print(f"workbook rows            {len(rows)}")
    print(f"entries written          {kept}")
    print(f"  positives              {stats['side_positives']}")
    print(f"  carveouts              {stats['side_carveouts']}")
    print("\nseverity (per entry):")
    for sev in ("warning", "cao", "nghiem_trong", "tich_cuc"):
        print(f"  {sev:14s} {stats['sev_' + sev]}")
    print("\nids:")
    print(f"  from the sheet         {stats['id_given']}")
    print(f"  recovered by content   {stats['id_recovered']}")
    print(f"  newly stamped          {stats['id_new']}")
    if stats["id_unknown"]:
        print(f"  in sheet, not in corpus {stats['id_unknown']}")
    dropped = len(by_id) - stats["id_given"] - stats["id_recovered"]
    print(f"\ndropped (absent from the workbook) {dropped}")
    print(f"groups                   {len(groups)}")
    print(f"needs review             {len(review)}")

    if a.dry_run:
        print("\ndry run — nothing written")
        return

    if not BACKUP.exists():
        shutil.copy2(corpus_path, BACKUP)
        print(f"\nbacked up -> {BACKUP.name}")
    else:
        # Second run. The backup is the pre-import corpus and must stay
        # that way, or the reference above stops meaning anything.
        print(f"\nkept      -> {BACKUP.name} (pre-import reference, untouched)")
    corpus_path.write_text(text, encoding="utf-8")
    print(f"wrote     -> {corpus_path.name}  ({len(text.splitlines())} lines)")

    header = "row\tissue\tid\toriginal\tconverted\tfragments"
    REVIEW.write_text("\n".join([header] + review), encoding="utf-8")
    print(f"review    -> {REVIEW.name}  ({len(review)} rows)")


if __name__ == "__main__":
    raise SystemExit(main())
