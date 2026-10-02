"""The rough tools in `tools/`: each runs as a script, refuses what it
should, and says why. Offline — JSON and xlsx in, xlsx or a number out."""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from tools import corpus_export, corpus_import, corpus_stamp_ids, qc_compare, qc_prepare, scan_report

from tests._paths import ROOT


@pytest.mark.parametrize("tool", [qc_prepare, qc_compare, scan_report])
def test_a_tool_missing_its_arguments_exits_2(tool):
    with pytest.raises(SystemExit) as e:
        tool.main([])
    assert e.value.code == 2


def test_the_tools_find_the_repo_from_their_new_home():
    assert corpus_import.CORPUS == ROOT / "knowledge/sentiment_agent/corpus.yaml"
    assert corpus_stamp_ids.CORPUS == corpus_import.CORPUS == corpus_export.CORPUS_PATH
    assert corpus_import.CORPUS.is_file()


# ── scan_report / corpus_export — files in, files out ───────────


def test_scan_report_writes_the_workbook(tmp_path):
    export = tmp_path / "export.json"
    export.write_text(json.dumps({"data": [{"call_id": "1", "call_scoring": {"Sentiment": [
        {"Criteria": "Thái độ ĐTV", "Result": "Tích cực", "traces": {"scanner": {"violation": False}}}]}}]}),
        encoding="utf-8")
    assert scan_report.main(["--input", str(export), "--output", str(tmp_path / "r.xlsx")]) == 0
    assert {"calls", "flagged", "summary"} <= set(load_workbook(tmp_path / "r.xlsx").sheetnames)


def test_corpus_export_has_one_row_per_entry_and_a_filter(tmp_path):
    import yaml

    assert corpus_export.main(["--output", str(tmp_path / "c.xlsx")]) == 0
    ws = load_workbook(tmp_path / "c.xlsx").active
    corpus = yaml.safe_load(corpus_export.CORPUS_PATH.read_text(encoding="utf-8"))
    entries = sum(len(g.get(k) or []) for g in corpus.values() if isinstance(g, dict)
                  for k in ("positives", "carveouts"))
    assert ws.max_row - 1 == entries and ws.auto_filter.ref






# ── qc_compare ──────────────────────────────────────────────────────────


def _qc_book(path: Path, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(["STT", "Call ID", "KQ Tự động", "KQ thủ công"])
    for r in rows:
        ws.append(r)
    wb.save(path)


def _scored(d: Path, stem: str, result: str):
    d.mkdir(parents=True, exist_ok=True)
    (d / f"{stem}.json").write_text(json.dumps({"Sentiment": [{"Criteria": "Thái độ ĐTV", "Result": result}]},
                                               ensure_ascii=False), encoding="utf-8")


def _compare(tmp_path):
    return qc_compare.main(["--scored", str(tmp_path / "out"), "--xlsx", str(tmp_path / "qc.xlsx"),
                            "--out", str(tmp_path / "cmp.xlsx")])


def test_qc_compare_joins_on_call_id_and_counts_agreement(tmp_path, capsys):
    _scored(tmp_path / "out", "E_x_CLID_0909_111", "Tích cực")
    _scored(tmp_path / "out", "E_y_CLID_0909_222", "Thái độ cao")
    _scored(tmp_path / "out", "E_z_CLID_0909_333", "Tích cực")  # not in QC's sheet
    _qc_book(tmp_path / "qc.xlsx", [[1, 111, "Tích cực", "Tích cực"], [2, 222, "Tích cực", "Tích cực"],
                                    [3, 444, "", "Tích cực"]])
    assert _compare(tmp_path) == 0
    out = capsys.readouterr().out
    assert "joined 2" in out and "only ours 1" in out and "only QC 1" in out and "1/2" in out
    assert {"calls", "manual_vs_ours", "auto_vs_ours"} <= set(load_workbook(tmp_path / "cmp.xlsx").sheetnames)


def test_qc_compare_says_when_there_is_no_ground_truth(tmp_path, capsys):
    _scored(tmp_path / "out", "E_x_111", "Tích cực")
    _qc_book(tmp_path / "qc.xlsx", [[1, 111, "Tích cực", None]])
    _compare(tmp_path)
    assert "NO GROUND TRUTH" in capsys.readouterr().out


def test_qc_compare_refuses_a_sheet_without_call_id(tmp_path):
    wb = Workbook()
    wb.active.append(["STT", "Tên"])
    wb.save(tmp_path / "qc.xlsx")
    (tmp_path / "out").mkdir()
    with pytest.raises(ValueError, match="Call ID"):
        _compare(tmp_path)


# ── corpus_stamp_ids / corpus_import ────────────────────────────────────


def test_stamp_ids_wants_exactly_one_of_apply_or_dry_run():
    with pytest.raises(SystemExit, match="exactly one"):
        corpus_stamp_ids.main([])


def test_corpus_import_writes_under_knowledge():
    assert "knowledge" in corpus_import.CORPUS.parts and ".prompts" not in str(corpus_import.CORPUS)

